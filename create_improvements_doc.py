import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(val='single', sz='4', color='3E484F')
    """
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_elm = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val","single")}" w:sz="{edge_data.get("sz","4")}" w:space="0" w:color="{edge_data.get("color","auto")}"/>')
            tcBorders.append(b_elm)
    tcPr.append(tcBorders)

def create_document():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(45, 55, 72)

    # Document Header / Hero Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("STITCH HIRING PIPELINE DIAGNOSTICS")
    title_run.font.name = 'Segoe UI'
    title_run.font.size = Pt(11)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(14, 116, 144) # Cyan-700

    main_title = doc.add_paragraph()
    main_title.paragraph_format.space_before = Pt(0)
    main_title.paragraph_format.space_after = Pt(8)
    main_run = main_title.add_run("100 Comprehensive Front-End Architectural & UX Improvements")
    main_run.font.name = 'Segoe UI'
    main_run.font.size = Pt(20)
    main_run.font.bold = True
    main_run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(16)
    meta_run = meta_p.add_run("Executive Engineering Blueprint | Scope: UI/UX, Performance, A11y, State, Visualization & Enterprise Readiness")
    meta_run.font.size = Pt(9.5)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    # Summary Callout Box
    summary_tbl = doc.add_table(rows=1, cols=1)
    summary_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_cell = summary_tbl.cell(0, 0)
    summary_cell.width = Inches(6.9)
    set_cell_background(summary_cell, "F0F9FF")
    set_cell_margins(summary_cell, top=140, bottom=140, left=180, right=180)
    set_cell_border(summary_cell, left=dict(val='single', sz='24', color='0284C7'),
                                 top=dict(val='single', sz='4', color='BAE6FD'),
                                 bottom=dict(val='single', sz='4', color='BAE6FD'),
                                 right=dict(val='single', sz='4', color='BAE6FD'))
    
    sp = summary_cell.paragraphs[0]
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    s_title = sp.add_run("Executive Summary")
    s_title.font.bold = True
    s_title.font.size = Pt(11)
    s_title.font.color.rgb = RGBColor(3, 105, 161)
    
    s_desc = summary_cell.add_paragraph()
    s_desc.paragraph_format.space_after = Pt(0)
    s_desc.paragraph_format.line_spacing = 1.15
    s_desc_run = s_desc.add_run(
        "This strategic document outlines 100 targeted frontend enhancements for the Stitch Hiring Pipeline Diagnostics platform. "
        "The recommendations are structured across 10 critical operational dimensions—spanning design systems, charting performance, "
        "data-grid virtualisation, bottleneck intelligence, keyboard accessibility, optimistic caching, and mobile responsiveness—to transform "
        "the application into an enterprise-grade, real-time recruitment command center."
    )
    s_desc_run.font.size = Pt(9.5)
    s_desc_run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    categories = [
        {
            "num": 1,
            "title": "Design System, Theming & Visual Aesthetics",
            "tag": "Aesthetics & HUD",
            "items": [
                ("Light / Dark / High-Contrast Theme Switcher", "Implement dynamic CSS variable toggle with prefers-color-scheme auto-detection and persistent localStorage state.", "High", "Low"),
                ("Semantic Severity Color Tokens", "Define standardized design tokens for critical, warning, healthy, and blocked states to eliminate hardcoded hex codes.", "High", "Low"),
                ("Consistent Spacing & Radius Scale", "Standardize border radius (rounded-md, rounded-xl) and padding scales across modals, cards, and data widgets.", "Medium", "Low"),
                ("Enhanced Glassmorphism & Cyber-HUD Depth", "Refine grid-bg-overlay with subtle radial vignette gradients and customizable backdrop-blur panels.", "Medium", "Medium"),
                ("Modern Icon System Optimization", "Audit and standardize Lucide icons using a centralized wrapper with consistent stroke widths (1.5px/2px) and dimensions.", "Medium", "Low"),
                ("Dynamic Status Badges & Glows", "Add pulse micro-animations (animate-ping) to active requisitions and high-urgency bottleneck alerts.", "Medium", "Low"),
                ("Skeleton Loading Shimmer Effects", "Replace generic spinner overlays with structured CSS shimmer skeletons matching Funnel and Table layouts.", "High", "Medium"),
                ("Subtle Micro-Animations", "Implement smooth hover transitions (ease-out), card elevation offsets (-translate-y-0.5), and interactive button ripples.", "Medium", "Low"),
                ("Typography Hierarchy Polish", "Fine-tune contrast ratios and line-heights between headings, metric values (font-mono), and secondary metadata.", "High", "Low"),
                ("Custom Cyber Focus Rings", "Add accessible and aesthetically unified cyber-focus outlines for keyboard navigation across all buttons and inputs.", "High", "Low"),
            ]
        },
        {
            "num": 2,
            "title": "Data Visualization & Recharts Analytics",
            "tag": "Charts & Metrics",
            "items": [
                ("Custom Interactive Tooltips in Funnel Charts", "Display stage conversion rate, drop-off delta vs. industry benchmark, and average days in stage on hover.", "High", "Medium"),
                ("Dynamic Chart Resizing with Debounce", "Prevent chart clipping or re-render thrashing during window resize events via debounced ResponsiveContainer.", "High", "Low"),
                ("Funnel Stage Zoom & Time Scrubber", "Add a brush/slider component to zoom into specific date ranges (Last 7d, 30d, 90d, All-time).", "High", "Medium"),
                ("Drop-Off Waterfall Visualization", "Add a dedicated Sankey/Waterfall diagram showing candidate dropout distribution across screening and technical stages.", "High", "High"),
                ("Cohort Analysis Heatmap", "Create a grid chart visualizing candidate pass-through velocity grouped by source channel and calendar week.", "High", "High"),
                ("Chart Export to PNG / SVG / PDF", "Integrate client-side canvas snapshotting (html-to-image) for downloading presentation-ready chart graphics.", "Medium", "Medium"),
                ("Empty State & Zero-Data Chart Fallbacks", "Provide sleek illustrative zero-data states with helpful diagnostic guidance rather than blank SVG canvases.", "Medium", "Low"),
                ("Synchronized Cross-Chart Hover Cursor", "Sync hover indicators across multiple timeline charts to inspect hiring velocity and volume simultaneously.", "Medium", "Medium"),
                ("Benchmark Overlay Lines", "Overlay target SLA durations and expected stage conversion rates as dashed reference lines (ReferenceLine).", "Medium", "Low"),
                ("Smooth Chart Entry Animations", "Configure custom spring animations on chart initial load (isAnimationActive, animationDuration=800).", "Low", "Low"),
            ]
        },
        {
            "num": 3,
            "title": "Diagnostic Funnel & Bottleneck Intelligence",
            "tag": "Diagnostic Intelligence",
            "items": [
                ("One-Click Diagnostic Filter Drill-Down", "Clicking a bottleneck stage in FunnelDashboard navigates to CandidateExplorer pre-filtered to stalled candidates.", "Critical", "Medium"),
                ("Automated Root-Cause Recommendation Badges", "Display smart tags (e.g., 'Interviewer Overload', 'High Drop-off at Tech Test', 'Slow Feedback Cycle').", "High", "Medium"),
                ("SLA Breach Threshold Indicators", "Add color-coded warning chips for stages where median time-in-stage exceeds configured SLA by >20%.", "High", "Low"),
                ("Interactive 'What-If' Funnel Simulator", "Provide slider controls allowing recruiters to simulate how fixing a stage drop-off impacts total hires.", "High", "High"),
                ("Multi-Job Pipeline Comparison Mode", "Add a split-screen or multi-select view to compare funnel health between 2-3 different roles side by side.", "High", "High"),
                ("Stage Velocity Breakdown (Median & P90)", "Display both median and 90th-percentile dwell time to highlight outlier candidates stalling the process.", "High", "Low"),
                ("Interviewer Capacity & Latency Gauges", "Add mini radial charts showing interviewer load and average feedback turnaround time per department.", "Medium", "Medium"),
                ("Stage Reordering & Custom Stage Mapping", "Allow dragging stages in the funnel view to customize evaluation flows per job requisition.", "Medium", "High"),
                ("Pipeline Velocity Trend Sparklines", "Embed mini sparkline charts inside stage summary headers showing 30-day velocity trajectory.", "Medium", "Low"),
                ("Diagnostic Alert Dismissal & Snoozing", "Allow hiring managers to acknowledge or snooze pipeline bottleneck alerts with custom notes.", "Medium", "Medium"),
            ]
        },
        {
            "num": 4,
            "title": "Candidate Explorer & Data Grid Capabilities",
            "tag": "Data Table & Operations",
            "items": [
                ("Virtual Scrolling / Windowing", "Implement @tanstack/react-virtual in CandidateExplorer to handle 10,000+ candidates smoothly at 60 FPS.", "Critical", "Medium"),
                ("Customizable & Reorderable Column Headers", "Allow users to toggle visibility and reorder table columns (Score, Source, Days in Stage, Recruiter, Actions).", "High", "Medium"),
                ("Multi-Column Sorting", "Support primary and secondary sorting (e.g., Sort by Stage then by Days in Stage DESC).", "High", "Low"),
                ("Persistent Filter Presets ('Saved Views')", "Allow recruiters to save complex filter configurations (e.g., 'Engineering - Stuck > 7 Days') in localStorage.", "High", "Medium"),
                ("Multi-Select Batch Actions Toolbar", "Add a floating sticky bottom action bar when candidates are selected (Batch Email, Change Stage, Reject, Export).", "Critical", "Medium"),
                ("Inline Quick-Edit Stage & Status", "Allow changing a candidate's status or assigned recruiter directly from a dropdown inside the table cell.", "High", "Low"),
                ("Global Fuzzy Search with Highlighted Matches", "Add debounced fuzzy search across candidate name, email, skills, and current company with keyword highlights.", "High", "Medium"),
                ("Candidate Quick-View Drawer / Side Panel", "Clicking a candidate opens a slide-over panel showing full activity timeline and interview feedback without losing scroll.", "High", "Medium"),
                ("Custom Badges for Candidate Sources", "Visually distinct colored chips for high-priority referral sources and diverse sourcing channels.", "Medium", "Low"),
                ("Copy-to-Clipboard Quick Actions", "One-click copy buttons for candidate emails, phone numbers, and profile share links with instant toast feedback.", "Low", "Low"),
            ]
        },
        {
            "num": 5,
            "title": "Modals, Forms & Requisition Creation",
            "tag": "Forms & Workflows",
            "items": [
                ("Multi-Step Form Progress Indicator", "Enhance NewRequisitionModal with a stepper (General Info -> Stages & SLA -> Hiring Team -> Budget).", "High", "Medium"),
                ("Form Validation with Zod & React Hook Form", "Replace manual form state with react-hook-form + zod for real-time, field-level validation and error messages.", "Critical", "Medium"),
                ("Auto-Save Drafts for New Requisitions", "Automatically persist unfinished requisition forms to sessionStorage to prevent data loss on accidental modal dismiss.", "High", "Low"),
                ("Dynamic Stage & SLA Rule Builder", "Allow adding, deleting, and reordering interview stages with custom SLA duration sliders during requisition creation.", "High", "Medium"),
                ("Department & Hiring Manager Typeahead Search", "Implement async search dropdown with avatar pills for assigning team members to requisitions.", "Medium", "Medium"),
                ("Modal Backdrop Blur & Smooth Scale-In Animation", "Replace abrupt modal popups with smooth spring transitions using Radix Dialog / Framer Motion.", "Medium", "Low"),
                ("Unsaved Changes Warning Dialog", "Trigger a confirmation alert when users attempt to close a modal with dirty form inputs.", "High", "Low"),
                ("Keyboard Accessibility for Modals", "Support Escape key close, initial input focus, and complete Tab focus trapping via Radix UI primitives.", "Critical", "Low"),
                ("Salary Range Slider with Currency Selector", "Replace plain text fields with dual-handle range sliders and multi-currency formatting ($ USD, € EUR, £ GBP).", "Medium", "Low"),
                ("Template Selector for Quick Requisitions", "Provide pre-built templates (e.g., 'Senior Frontend Engineer', 'Product Manager') to auto-populate stages and SLAs.", "Medium", "Low"),
            ]
        },
        {
            "num": 6,
            "title": "CSV Import & Bulk Data Processing",
            "tag": "Data Ingestion",
            "items": [
                ("Interactive Column Mapping Screen", "In CSVImportModal, display a visual column-to-field mapper with dropdowns and live 3-row data previews.", "Critical", "Medium"),
                ("Drag-and-Drop File Zone with Visual States", "Highlight active drop zones with border glow, file icon animation, and clear size/format constraints.", "High", "Low"),
                ("Pre-Import Data Validation & Error Highlighting", "Validate rows before sending to backend, displaying invalid rows (e.g., bad emails, duplicates) in a review grid.", "Critical", "Medium"),
                ("Import Progress Bar with Real-Time Percentage & ETA", "Display real-time streaming progress when importing large candidate rosters.", "High", "Medium"),
                ("Downloadable Sample CSV Template", "Add a prominent button to download pre-formatted .csv and .xlsx templates with demo data.", "High", "Low"),
                ("Partial Import Resolution & Retry Options", "Allow skipping invalid rows and importing valid ones, with a downloadable CSV of failed rows and error reasons.", "High", "Medium"),
                ("Duplicate Detection Alert", "Preview detected duplicate candidate emails before committing the import with options to Overwrite, Skip, or Merge.", "High", "Medium"),
                ("Import History & Audit Log", "Maintain a table in Settings showing past CSV imports, who uploaded them, timestamp, and row count.", "Medium", "Medium"),
                ("Support for Multiple Date & Number Formats", "Automatically parse varying date formats (DD/MM/YYYY, MM/DD/YYYY, ISO-8601) gracefully.", "Medium", "Low"),
                ("Direct ATS Integration Connectors Mockup/UI", "Add tabs in import modal for 1-click sync from Greenhouse, Lever, Workday, and BambooHR.", "Medium", "High"),
            ]
        },
        {
            "num": 7,
            "title": "Performance, Caching & Web Vitals",
            "tag": "Performance & Architecture",
            "items": [
                ("Client-Side Query Caching (TanStack Query / SWR)", "Replace raw fetch calls in AppContext with cached, deduplicated hooks with automatic background revalidation.", "Critical", "Medium"),
                ("Code Splitting & Lazy Loading for Heavy Components", "Dynamic import (next/dynamic) for Recharts graphs, CSV parser (papaparse), and heavy modals to cut bundle size.", "High", "Low"),
                ("Font & Asset Optimization", "Ensure Google Fonts (Inter, JetBrains Mono) are loaded via next/font/google with display: swap to eliminate CLS.", "High", "Low"),
                ("Debounced API Requests", "Apply 300ms debounce on global search inputs, date range changes, and slider adjustments to prevent API spam.", "High", "Low"),
                ("Optimistic UI Updates", "Immediately reflect stage changes, status updates, and note additions in the UI before network confirmation.", "High", "Medium"),
                ("Memoization of Heavy Calculations", "Wrap complex funnel metric aggregations and bottleneck calculations with useMemo to prevent UI re-render lag.", "High", "Low"),
                ("Image & Avatar Optimization", "Use next/image with WebP/AVIF formats and placeholder blur for recruiter and candidate profile pictures.", "Medium", "Low"),
                ("Service Worker / Offline Readiness", "Cache core UI shell and read-only pipeline snapshots for offline inspection during poor network connectivity.", "Medium", "High"),
                ("Tree-Shaking Icons", "Import Lucide icons as named direct imports to avoid pulling the entire icon library into the client bundle.", "Medium", "Low"),
                ("Core Web Vitals Monitoring", "Integrate lightweight performance reporting (web-vitals) logging LCP, FID, and CLS scores to analytics.", "Low", "Low"),
            ]
        },
        {
            "num": 8,
            "title": "Accessibility & Usability (a11y)",
            "tag": "WCAG 2.1 AA Compliance",
            "items": [
                ("WCAG 2.1 AA Color Contrast Verification", "Ensure all muted text against dark backgrounds (#0b1326) exceeds the minimum 4.5:1 contrast ratio.", "Critical", "Low"),
                ("Full ARIA Labels & Dynamic Live Regions", "Add aria-expanded, aria-controls, aria-haspopup, and aria-live='polite' to notification bells and filters.", "High", "Low"),
                ("Global Keyboard Shortcut System", "Add global shortcuts (/ search, N new requisition, C candidates, Esc close) with an interactive '?' modal.", "High", "Medium"),
                ("Screen Reader Accessible Data Tables", "Provide semantic <caption>, <th scope='col'>, and descriptive sr-only summary tags for complex funnel tables.", "High", "Low"),
                ("Reduced Motion Support", "Wrap CSS animations with @media (prefers-reduced-motion: reduce) to disable intense transitions.", "Medium", "Low"),
                ("Visible Focus Order & Skip-to-Content Link", "Add a top-level skip link for keyboard users to bypass navigation straight to the main analytics content.", "Medium", "Low"),
                ("Form Accessibility & Error Association", "Ensure all inputs have associated <label> tags with htmlFor and error states tied via aria-describedby.", "High", "Low"),
                ("Tooltips on Truncated Text", "Automatically display full text tooltips when candidate names, job titles, or emails are truncated with ellipsis.", "Medium", "Low"),
                ("Semantic HTML Structure", "Refactor generic <div> containers into semantic <header>, <nav>, <main>, <aside>, and <section> tags.", "High", "Low"),
                ("Color-Blind Safe Chart Modes", "Provide alternative chart palettes using patterns or distinct color-blind safe palettes (Viridis / Okabe-Ito).", "Medium", "Medium"),
            ]
        },
        {
            "num": 9,
            "title": "User Feedback, Notifications & Micro-Interactions",
            "tag": "Feedback & Interactions",
            "items": [
                ("Toast Notification System", "Replace native browser alerts with rich stacked toast notifications (Sonner / Radix) with 'Undo' actions.", "Critical", "Low"),
                ("Interactive Notification Drawer", "Enhance NotificationDropdown with filter tabs (All, Bottlenecks, Mentions, SLA), 'Mark all as read', and links.", "High", "Medium"),
                ("Auditory Feedback Option", "Optional subtle, professional audio cues for high-priority pipeline breach alerts and successful imports.", "Low", "Low"),
                ("Inline Confirmation Popovers", "Replace invasive browser confirms with sleek inline popovers ('Are you sure?') with Confirm/Cancel buttons.", "High", "Low"),
                ("Live Recruiter Collaboration Avatars", "Show 'who is currently viewing this job/candidate' avatar circles in headers using WebSocket presence.", "Medium", "High"),
                ("Copyable Deep Links", "Add a 'Share View' button that serializes active filters, date ranges, and job IDs directly into URL parameters.", "High", "Low"),
                ("Interactive Empty States with Action CTAs", "Provide clear calls to action when no data matches (e.g., 'No candidates stuck — [Clear Filters]').", "High", "Low"),
                ("Breadcrumb Navigation", "Add breadcrumbs across nested pages (e.g., Jobs > Senior Frontend Engineer > Funnel Diagnostics).", "Medium", "Low"),
                ("Progressive Disclosure for Advanced Metrics", "Provide expandable accordion cards for deep statistical variance and standard deviations.", "Medium", "Low"),
                ("Interactive Help & Onboarding Tour", "Add a step-by-step guided product tour highlighting key diagnostic HUD features for new recruiters.", "Medium", "Medium"),
            ]
        },
        {
            "num": 10,
            "title": "Responsive Layout, Settings & Enterprise Controls",
            "tag": "Enterprise & Layout",
            "items": [
                ("Mobile Bottom Navigation Bar Enhancement", "Polish BottomNav with active indicator pill animations, badge counters, and thumb-friendly touch targets.", "High", "Low"),
                ("Collapsible & Resizable Sidebar", "Allow recruiters on smaller screens to collapse Sidebar into an icon-only rail or drag to resize panel width.", "High", "Medium"),
                ("Responsive Table Card View for Mobile", "On mobile viewports (<768px), automatically transform wide candidate tables into scannable stacked cards.", "High", "Medium"),
                ("Executive PDF One-Pager Generator", "Build a dedicated print/PDF stylesheet (@media print) formatting Executive Overview into a clean summary document.", "High", "Medium"),
                ("Configurable SLA Rule Engine in Settings", "In SettingsView, allow setting custom SLA thresholds (e.g., Resume Review: 2d, Tech: 5d) per department.", "High", "Medium"),
                ("Webhook & Slack Alert Configuration UI", "Add a settings panel to configure real-time Slack/Teams alerts when pipeline bottlenecks exceed thresholds.", "High", "Medium"),
                ("Role-Based View Customization (RBAC UI)", "Render distinct interface modes based on persona: Executive View, Recruiter View, and Hiring Manager View.", "High", "High"),
                ("Timezone & Regional Localization", "Provide user settings for date formats, currency symbols, and multi-timezone interview scheduling displays.", "Medium", "Medium"),
                ("Automated Error Boundary & Crash Recovery UI", "Implement React 19 ErrorBoundary wrappers around widgets with a 'Reload Component' recovery button.", "Critical", "Low"),
                ("Live Diagnostic Health Score Meter", "Add a prominent HUD widget calculating an overall 'Pipeline Health Index' (0-100%) aggregating velocity and SLA metrics.", "High", "Medium"),
            ]
        }
    ]

    total_counter = 0

    for cat in categories:
        # Category Heading
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(14)
        h_p.paragraph_format.space_after = Pt(4)
        h_p.paragraph_format.keep_with_next = True

        h_run_num = h_p.add_run(f"Section {cat['num']}: ")
        h_run_num.font.bold = True
        h_run_num.font.size = Pt(13)
        h_run_num.font.color.rgb = RGBColor(14, 116, 144)

        h_run_title = h_p.add_run(cat['title'])
        h_run_title.font.bold = True
        h_run_title.font.size = Pt(13)
        h_run_title.font.color.rgb = RGBColor(15, 23, 42)

        # Table for the 10 items
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        col_widths = [Inches(0.5), Inches(4.5), Inches(0.9), Inches(1.0)]

        # Header Row
        hdr_cells = table.rows[0].cells
        hdr_titles = ["#", "Improvement & Technical Recommendation", "Impact", "Effort"]
        
        for idx, cell in enumerate(hdr_cells):
            cell.width = col_widths[idx]
            set_cell_background(cell, "0F172A") # Slate-900
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if idx in (0, 2, 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(hdr_titles[idx])
            run.font.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for item in cat['items']:
            total_counter += 1
            row_cells = table.add_row().cells
            bg_color = "F8FAFC" if total_counter % 2 == 0 else "FFFFFF"

            for idx, cell in enumerate(row_cells):
                cell.width = col_widths[idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                set_cell_border(cell, bottom=dict(val='single', sz='2', color='E2E8F0'),
                                     top=dict(val='single', sz='2', color='E2E8F0'))

            # Cell 0: Number
            p0 = row_cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(str(total_counter))
            r0.font.bold = True
            r0.font.size = Pt(9)
            r0.font.color.rgb = RGBColor(71, 85, 105)

            # Cell 1: Improvement Name + Description
            p1 = row_cells[1].paragraphs[0]
            p1.paragraph_format.space_after = Pt(2)
            p1.paragraph_format.line_spacing = 1.1
            r1_title = p1.add_run(item[0] + "\n")
            r1_title.font.bold = True
            r1_title.font.size = Pt(9.5)
            r1_title.font.color.rgb = RGBColor(15, 23, 42)

            r1_desc = p1.add_run(item[1])
            r1_desc.font.size = Pt(8.5)
            r1_desc.font.color.rgb = RGBColor(71, 85, 105)

            # Cell 2: Impact
            p2 = row_cells[2].paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(item[2])
            r2.font.bold = True
            r2.font.size = Pt(8.5)
            if item[2] == "Critical":
                r2.font.color.rgb = RGBColor(220, 38, 38)
            elif item[2] == "High":
                r2.font.color.rgb = RGBColor(217, 119, 6)
            elif item[2] == "Medium":
                r2.font.color.rgb = RGBColor(14, 116, 144)
            else:
                r2.font.color.rgb = RGBColor(100, 116, 139)

            # Cell 3: Effort
            p3 = row_cells[3].paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.paragraph_format.space_after = Pt(0)
            r3 = p3.add_run(item[3])
            r3.font.size = Pt(8.5)
            r3.font.color.rgb = RGBColor(71, 85, 105)

        # Space after table
        sp_p = doc.add_paragraph()
        sp_p.paragraph_format.space_after = Pt(8)

    # Output filename
    output_path = os.path.join(os.getcwd(), "100_FrontEnd_Improvements_Stitch_Dashboard.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    create_document()
