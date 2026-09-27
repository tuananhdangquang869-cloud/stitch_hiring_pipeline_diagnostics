import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.color.rgb = RGBColor(11, 44, 90) # Navy
        run.font.size = Pt(16)
        run.bold = True
    elif level == 2:
        run.font.color.rgb = RGBColor(0, 114, 188) # Blue
        run.font.size = Pt(13)
        run.bold = True
    elif level == 3:
        run.font.color.rgb = RGBColor(50, 50, 50)
        run.font.size = Pt(11)
        run.bold = True
    return h

def generate_report():
    doc = docx.Document()
    
    # Page Setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    # Base directory for screenshots
    base_dir = r"m:\stitch_hiring_pipeline_diagnostics\stitch_hiring_pipeline_diagnostics"

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("BÁO CÁO PHÂN TÍCH KIẾN TRÚC GIAO DIỆN")
    t_run.font.size = Pt(20)
    t_run.bold = True
    t_run.font.color.rgb = RGBColor(11, 44, 90)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(16)
    s_run = sub_p.add_run("Dự án: Áp Lực Kế (Pressure Gauge) — B2B Recruitment Funnel Analytics")
    s_run.font.size = Pt(12)
    s_run.font.italic = True
    s_run.font.color.rgb = RGBColor(80, 80, 80)

    # Meta Info Box
    tbl_meta = doc.add_table(rows=1, cols=1)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_meta = tbl_meta.cell(0, 0)
    set_cell_background(cell_meta, "F0F4F8")
    set_cell_margins(cell_meta, top=140, bottom=140, left=200, right=200)
    p_meta = cell_meta.paragraphs[0]
    p_meta.paragraph_format.line_spacing = 1.2
    r = p_meta.add_run("📌 Tổng hợp kết quả phân tích: ")
    r.bold = True
    p_meta.add_run("Toàn bộ 18 thư mục xuất từ Google Stitch đã được kiểm tra chi tiết mã nguồn HTML, CSS Tailwind tokens và tệp ảnh chụp màn hình (screen.png). Dữ liệu thiết kế được chuẩn hóa thành 6 màn hình chức năng chính + bộ quy chuẩn nhận diện thương hiệu (Design System & Logo).\n")
    r2 = p_meta.add_run("🎯 Mục tiêu: ")
    r2.bold = True
    p_meta.add_run("Xác định biến thể hoàn thiện nhất (Production-ready pick) cho từng màn hình để tiến hành hiện thực hóa ứng dụng Web chuẩn B2B SaaS.")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. BẢNG TỔNG HỢP 6 MÀN HÌNH CHÍNH
    add_styled_heading(doc, "1. BẢNG TỔNG HỢP NHÓM 6 MÀN HÌNH CỐT LÕI", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["STT / Màn hình", "Các thư mục Stitch", "Phiên bản đề xuất (Pick)", "Đánh giá & Lý do chọn"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "0B2C5A")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)

    screens_summary = [
        (
            "1. Phân tích phễu tuyển dụng\n(Funnel Analytics & Pipeline Diagnostics)",
            "• main_dashboard_interactive_states\n• main_dashboard_funnel_diagnostics\n• main_dashboard_kinetic_theme\n• funnel_analytics_diagnostics\n• dashboard_loading_state",
            "main_dashboard_interactive_states\n(+ tích hợp chỉ số funnel_analytics)",
            "Layout 3 cột chuẩn chỉn: Cột chọn Job bên trái, đồ thị phễu tuyển dụng trực quan ở giữa, Drawer chẩn đoán và nhật ký rớt ứng viên (Leakage Log) bên phải."
        ),
        (
            "2. Chi tiết & Chẩn đoán giai đoạn\n(Stage Detail Modal / Inspector)",
            "• stage_detail_full_view\n• stage_detail_kinetic_theme",
            "stage_detail_full_view",
            "Modal chi tiết hoàn chỉnh: Thời gian lưu tại vòng (Avg Time), lý do rớt top đầu (Skill gap, Lương, Văn hóa), danh sách ứng viên kèm mã lỗi (ERR_TECH_GAP...)."
        ),
        (
            "3. Danh sách ứng viên\n(Candidate Explorer)",
            "• candidate_explorer",
            "candidate_explorer",
            "Bảng quản trị ứng viên chuyên sâu: Lọc theo vòng tuyển, thanh tìm kiếm, đo chỉ số sức khỏe hồ sơ (Health Index 0-100), cảnh báo SLA thời gian và nút Diagnose."
        ),
        (
            "4. Quản trị vị trí tuyển dụng\n(Job Management Hub)",
            "• job_selector_expanded_state",
            "job_selector_expanded_state",
            "Màn hình quản lý Job Requisitions: Danh sách job đang tuyển, số lượng ứng tuyển, tỷ lệ chuyển đổi, chỉ số Time-to-fill, Funnel Health Gauge và nút thao tác nhanh."
        ),
        (
            "5. Nạp dữ liệu ứng viên CSV\n(CSV Data Import Modal)",
            "• csv_import_modal",
            "csv_import_modal",
            "Modal tải lên dữ liệu CSV: Vùng kéo thả file, tự động khớp trường dữ liệu (Header Mapping), bảng xem trước dữ liệu mẫu và nút xác nhận Import hàng loạt."
        ),
        (
            "6. Tổng quan điều hành\n(Executive Overview & Mobile Summary)",
            "• mobile_summary_dashboard\n• mobile_summary_view_detailed\n• mobile_summary_kinetic_theme\n• executive_overview\n• dashboard_empty_state",
            "mobile_summary_dashboard\n(Chuyển đổi Responsive Web)",
            "Màn hình duy nhất có đầy đủ KPI điều hành cấp cao: Cảnh báo nghẽn nghiêm trọng (Critical Bottleneck), quá tải Recruiter, chỉ số SLA, biểu đồ nhận việc (Acceptance Rate)."
        )
    ]

    for idx, (scr, flds, pick, reason) in enumerate(screens_summary):
        row = table.add_row()
        cells = row.cells
        bg_color = "FFFFFF" if idx % 2 == 0 else "F7FAFC"
        for i, text in enumerate([scr, flds, pick, reason]):
            cells[i].text = text
            set_cell_background(cells[i], bg_color)
            set_cell_margins(cells[i], top=100, bottom=100, left=100, right=100)
            p = cells[i].paragraphs[0]
            p.runs[0].font.size = Pt(9.5)
            if i == 2:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(0, 102, 160)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 2. CHI TIẾT TỪNG MÀN HÌNH KÈM HÌNH ẢNH MINH HỌA
    add_styled_heading(doc, "2. CHI TIẾT TỪNG MÀN HÌNH VÀ PHIÊN BẢN ĐỀ XUẤT", level=1)

    screen_details = [
        {
            "num": "Màn hình 1",
            "title": "Phân tích Phễu Tuyển dụng (Funnel Analytics & Pipeline Diagnostics)",
            "folders": "main_dashboard_interactive_states, main_dashboard_funnel_diagnostics, main_dashboard_kinetic_theme, funnel_analytics_diagnostics, dashboard_loading_state",
            "pick": "main_dashboard_interactive_states (kết hợp các chỉ số chuyên sâu từ funnel_analytics_diagnostics)",
            "desc": "Đây là màn hình cốt lõi nhất của sản phẩm 'Áp Lực Kế', giúp giải quyết bài toán: 'Ứng viên đang rớt ở đâu và tại sao?'. Giao diện gồm 3 thành phần chính:\n"
                    "1. Thanh điều hướng & Bộ chọn vị trí (Sidebar & Job Selector): Cho phép chuyển đổi nhanh giữa các Requisition (Sr. Backend Engineer, Product Manager, UX Designer...).\n"
                    "2. Phễu chuyển đổi trung tâm (Funnel Visualization): Thể hiện trực quan số lượng ứng viên qua từng vòng (Applied -> Screened -> Tech Interview -> Offer -> Hired) kèm tỷ lệ rớt (Leakage rate).\n"
                    "3. Drawer chẩn đoán chuyên sâu (Stage Diagnostics Inspector): Mở ra khi click vào bất kỳ giai đoạn nào, hiển thị tỷ lệ lý do rớt (Failed Tech Screen, Salary, Ghosted) và danh sách ứng viên rớt gần nhất.",
            "img": os.path.join(base_dir, "main_dashboard_interactive_states", "screen.png")
        },
        {
            "num": "Màn hình 2",
            "title": "Chi tiết Chẩn đoán Giai đoạn (Stage Detail Modal / Inspector)",
            "folders": "stage_detail_full_view, stage_detail_kinetic_theme",
            "pick": "stage_detail_full_view",
            "desc": "Màn hình dạng Modal phóng to để HR Manager xem phân tích chi tiết một vòng cụ thể (ví dụ: Tech Interview):\n"
                    "• Thẻ chỉ số: Thời gian lưu hồ sơ trung bình (12.4 days) và Tổng lượng ứng viên rớt (142 candidates).\n"
                    "• Biểu đồ phân rã nguyên nhân: Technical Skill Gap (42%), Salary Mismatch (28%), Cultural Fit (18%), Candidate Withdrew (12%).\n"
                    "• Bảng danh sách chi tiết: ID ứng viên kèm mã định danh lỗi (ERR_TECH_GAP, ERR_SALARY, ERR_CULTURE) và số ngày tồn đọng.",
            "img": os.path.join(base_dir, "stage_detail_full_view", "screen.png")
        },
        {
            "num": "Màn hình 3",
            "title": "Danh sách Quản lý Ứng viên (Candidate Explorer)",
            "folders": "candidate_explorer",
            "pick": "candidate_explorer",
            "desc": "Màn hình quản trị dữ liệu chi tiết toàn bộ hồ sơ trong đường ống tuyển dụng:\n"
                    "• Thanh bộ lọc theo trạng thái: All Stages, Screening, Technical, Offer.\n"
                    "• Bảng dữ liệu ứng viên: Hiển thị Candidate ID, Họ tên, Vòng tuyển dụng hiện tại, Thời gian đã ở trong vòng (có icon cảnh báo SLA quá hạn), Recruiter phụ trách.\n"
                    "• Điểm sức khỏe hồ sơ (Health Index 0-100): Cột đo trực quan bằng thanh progress bar phân màu (Xanh lá = tốt, Vàng = rủi ro, Đỏ = khẩn cấp).\n"
                    "• Thao tác nhanh: Nút 'Diagnose' mở ngay chẩn đoán nguyên nhân hồ sơ bị tắc.",
            "img": os.path.join(base_dir, "candidate_explorer", "screen.png")
        },
        {
            "num": "Màn hình 4",
            "title": "Quản trị Vị trí Tuyển dụng (Job Management Hub)",
            "folders": "job_selector_expanded_state",
            "pick": "job_selector_expanded_state",
            "desc": "Màn hình dành riêng cho việc theo dõi danh mục các vị trí đang mở tuyển dụng (Requisitions):\n"
                    "• Cột danh sách Requisition: Hiển thị mã vị trí, phòng ban, số lượng hồ sơ hiện có và xu hướng chuyển đổi (+4.2%, -1.1%).\n"
                    "• Thẻ tổng quan hiệu suất: Total Applicants, Time to Fill (EST) với cảnh báo SLA, Funnel Health (Good / Needs Attention).\n"
                    "• Thao tác Requisition: Nút tạo mới 'New Job Requisition', Chỉnh sửa 'Edit', hoặc Tạm dừng 'Suspend'.",
            "img": os.path.join(base_dir, "job_selector_expanded_state", "screen.png")
        },
        {
            "num": "Màn hình 5",
            "title": "Nạp dữ liệu ứng viên từ CSV (CSV Data Import Modal)",
            "folders": "csv_import_modal",
            "pick": "csv_import_modal",
            "desc": "Màn hình Modal hỗ trợ người dùng đưa dữ liệu ứng viên thực tế từ các nguồn ATS bên ngoài vào hệ thống:\n"
                    "• Vùng tải tệp: Kéo thả file .csv, .xlsx dung lượng tối đa 5MB.\n"
                    "• Trình ánh xạ cột thông minh (Header Mapping): Tự động khớp các cột trong CSV (Full Name, Email Addr, Position, Status) với các trường hệ thống tương ứng.\n"
                    "• Bảng xem trước dữ liệu (Data Mapping Preview): Xem trước 5 dòng đầu đã được parse với badge trạng thái đẹp mắt trước khi bấm 'Import'.",
            "img": os.path.join(base_dir, "csv_import_modal", "screen.png")
        },
        {
            "num": "Màn hình 6",
            "title": "Tổng quan Điều hành Cấp cao (Executive Overview & Summary)",
            "folders": "mobile_summary_dashboard, mobile_summary_view_detailed, mobile_summary_kinetic_theme, executive_overview, dashboard_empty_state",
            "pick": "mobile_summary_dashboard (sử dụng bố cục và dữ liệu chỉ số để dựng bản Executive Overview responsive)",
            "desc": "Màn hình tổng hợp dành cho C-Level / Head of HR xem nhanh tình hình toàn công ty:\n"
                    "• Banner cảnh báo điểm nghẽn (Critical Bottleneck): Thông báo sụt giảm pass rate bất thường.\n"
                    "• Cảnh báo công suất (Capacity Warning): Cảnh báo khi tải công việc của Recruiter vượt ngưỡng tối ưu 15%.\n"
                    "• Các chỉ số chiến lược: Active Reqs, Time to Fill, Offer Acceptance Rate (với biểu đồ Donut 88.5%).\n"
                    "• Tổng quan phễu lũy kế năm (Funnel Diagnostic YTD): Tỷ lệ rơi rớt qua các nấc Sourced -> Screened -> Interview.",
            "img": os.path.join(base_dir, "mobile_summary_dashboard", "screen.png")
        }
    ]

    for item in screen_details:
        add_styled_heading(doc, f"{item['num']}: {item['title']}", level=2)
        
        p_desc = doc.add_paragraph()
        r_f = p_desc.add_run("📁 Các thư mục liên quan: ")
        r_f.bold = True
        p_desc.add_run(item["folders"] + "\n")
        
        r_p = p_desc.add_run("⭐ Phiên bản đề xuất: ")
        r_p.bold = True
        r_p.font.color.rgb = RGBColor(0, 102, 160)
        p_desc.add_run(item["pick"] + "\n\n")
        
        p_desc.add_run(item["desc"])
        p_desc.paragraph_format.space_after = Pt(6)

        # Add image if exists
        if os.path.exists(item["img"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(12)
            try:
                # Set width to 5.5 inches for clean layout
                run_img = p_img.add_run()
                run_img.add_picture(item["img"], width=Inches(5.5))
                
                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(14)
                r_cap = p_cap.add_run(f"Hình minh họa: {item['title']}")
                r_cap.font.italic = True
                r_cap.font.size = Pt(9)
                r_cap.font.color.rgb = RGBColor(100, 100, 100)
            except Exception as e:
                p_img.add_run(f"[Không thể tải ảnh: {e}]")

    # 3. PHÂN TÍCH HỆ THỐNG THIẾT KẾ & THƯƠNG HIỆU
    add_styled_heading(doc, "3. PHÂN TÍCH DESIGN SYSTEM VÀ NHẬN DIỆN THƯƠNG HIỆU", level=1)
    
    p_ds = doc.add_paragraph()
    p_ds.add_run("Trong gói xuất dữ liệu từ Google Stitch có 2 bộ Design System và 1 tệp Logo thương hiệu:")

    tbl_ds = doc.add_table(rows=4, cols=3)
    tbl_ds.alignment = WD_TABLE_ALIGNMENT.CENTER
    ds_headers = ["Thành phần / Thư mục", "Đặc điểm nhận diện", "Đánh giá & Khuyến nghị áp dụng"]
    for i, h in enumerate(ds_headers):
        tbl_ds.cell(0, i).text = h
        set_cell_background(tbl_ds.cell(0, i), "0B2C5A")
        set_cell_margins(tbl_ds.cell(0, i), top=100, bottom=100, left=100, right=100)
        p = tbl_ds.cell(0, i).paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)

    ds_rows = [
        (
            "Pressure Gauge Design System\n(pressure_gauge/DESIGN.md)",
            "• Nền tối xanh navy sâu (#0B1326)\n• Màu chủ đạo: Cyan / Xanh biển sáng (#8ED5FF, #38BDF8)\n• Màu phụ: Hổ phách cảnh báo (#FFB95F), Đỏ san hô (#FFB4AB)\n• Typography: Inter & JetBrains Mono",
            "KHUYẾN NGHỊ LÀM GIAO DIỆN CHÍNH (Default Theme). Cho cảm giác công nghệ cao (Deep-tech / Cybersecurity style), độ tương phản cao, chuyên nghiệp cho dashboard phân tích."
        ),
        (
            "Kinetic Pressure Design System\n(kinetic_pressure/DESIGN.md)",
            "• Nền tối than ấm (#16171A)\n• Màu chủ đạo: Vàng cam hổ phách (#FFC174, #F59E0B)\n• Typography: Geist Font",
            "Tùy chọn Theme phụ (Secondary Theme). Có thể cung cấp dưới dạng tính năng đổi Theme (Theme Toggle) trong phần Cài đặt (Settings)."
        ),
        (
            "Logo Áp Lực Kế\n(p_l_c_k_logo/screen.png)",
            "Biểu tượng kim đo áp suất kết hợp hình phễu tuyển dụng với gam màu xanh Cyan hiện đại.",
            "Sử dụng làm Logo chuẩn ở góc trên bên trái Sidebar của toàn bộ hệ thống ứng dụng."
        )
    ]

    for idx, (name, spec, rec) in enumerate(ds_rows):
        row_cells = tbl_ds.rows[idx + 1].cells
        for i, text in enumerate([name, spec, rec]):
            row_cells[i].text = text
            set_cell_background(row_cells[i], "FFFFFF" if idx % 2 == 0 else "F7FAFC")
            set_cell_margins(row_cells[i], top=100, bottom=100, left=100, right=100)
            p = row_cells[i].paragraphs[0]
            p.runs[0].font.size = Pt(9.5)
            if i == 0:
                p.runs[0].font.bold = True

    # Logo image embedding
    logo_path = os.path.join(base_dir, "p_l_c_k_logo", "screen.png")
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(12)
        p_logo.paragraph_format.space_after = Pt(4)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.5))
        
        p_logocap = doc.add_paragraph()
        p_logocap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logocap.paragraph_format.space_after = Pt(12)
        r_lc = p_logocap.add_run("Logo nhận diện thương hiệu 'Áp Lực Kế'")
        r_lc.font.italic = True
        r_lc.font.size = Pt(9)
        r_lc.font.color.rgb = RGBColor(100, 100, 100)

    # 4. KẾ HOẠCH BƯỚC TIẾP THEO
    add_styled_heading(doc, "4. ĐỀ XUẤT LỘ TRÌNH THỰC THI (NEXT STEPS)", level=1)
    
    steps = [
        ("Bước 1: Xác nhận kiến trúc", "Người dùng duyệt qua danh sách 6 màn hình và phiên bản đề xuất trong báo cáo này."),
        ("Bước 2: Xây dựng Core Framework & Design Tokens", "Tạo cấu trúc dự án chuẩn (HTML5, Vanilla CSS, JS) với toàn bộ biến màu CSS tokens, font JetBrains Mono / Inter, và layout khung 3 cột responsive."),
        ("Bước 3: Tích hợp và liên kết 6 màn hình", "Lắp ghép 6 màn hình hoàn chỉnh, cho phép chuyển đổi mượt mà giữa Funnel Analytics, Candidate Explorer, Job Management, CSV Import Modal và Stage Diagnostics Drawer."),
        ("Bước 4: Bổ sung tính tương tác & dữ liệu mẫu", "Kích hoạt các tính năng tương tác thực tế: Click chọn Job để cập nhật phễu, Click vào vòng phễu để mở Drawer chẩn đoán, Lọc bảng ứng viên, Import CSV giả lập.")
    ]

    for st_title, st_desc in steps:
        p_step = doc.add_paragraph()
        r_st = p_step.add_run(f"• {st_title}: ")
        r_st.bold = True
        r_st.font.color.rgb = RGBColor(0, 51, 102)
        p_step.add_run(st_desc)

    # Save document
    output_path = r"m:\stitch_hiring_pipeline_diagnostics\Bao_Cao_Phan_Tich_Ap_Luc_Ke.docx"
    doc.save(output_path)
    print(f"Report successfully generated at: {output_path}")

if __name__ == "__main__":
    generate_report()
